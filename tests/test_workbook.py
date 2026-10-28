"""엑셀 워크북 첫 화면 API — 한 줄 요약, 등록 잔고로 리밸런싱(이번 달 납입), 보고서 엑셀. 임시 규칙·DB, 네트워크 없음."""
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient

from backend import storage
from backend.main import app

ACCOUNTS = """accounts:
  - id: me_wb
    label: 테스트 연금저축
    pool: me_wb_pool
    owner: me
    broker: mirae
    account_type: pension_savings
    pocket: pension
    active: true
pools:
  me_wb_pool:
    label: 테스트 연금 — 매수만
    strategy: all_weather
    mode: buy_only
"""
AMAP = """assets:
  KODEX 200:
    ticker: '069500'
    asset_class: stock
    role: core
    status: confirmed
class_defaults: {}
"""


@pytest.fixture
def client(tmp_path, monkeypatch):
    (tmp_path / "accounts.yaml").write_text(ACCOUNTS, encoding="utf-8")
    (tmp_path / "asset_map.yaml").write_text(AMAP, encoding="utf-8")
    monkeypatch.setenv("RETIRE_ACCOUNTS_PATH", str(tmp_path / "accounts.yaml"))
    monkeypatch.setenv("RETIRE_ASSET_MAP_PATH", str(tmp_path / "asset_map.yaml"))
    monkeypatch.setattr(storage, "DB_PATH", tmp_path / "ledger.sqlite3")
    with TestClient(app, base_url="http://127.0.0.1", headers={"Origin": "http://127.0.0.1", "X-Retire-Local": "1"}) as c:
        r = c.post("/api/register/commit", json={"account_id": "me_wb", "source": "file", "as_of": date.today().isoformat(), "screen_total": None,
                                                 "rows": [{"ticker_name": "KODEX 200", "ticker": "069500", "qty": 100, "price": 30000, "asset_class": "stock", "buy_value": 2_500_000}]})
        assert r.status_code == 200, r.text
        yield c


def test_summary_line_has_net_returns_and_rebalance_verdict(client):
    d = client.get("/api/workbook/summary").json()
    assert "순자산 300만원" in d["line"] and "불어난 돈 50만원 +20.00%" in d["line"]
    assert "리밸런싱 필요: 테스트 연금 주식 +70.0%p" in d["line"]
    assert d["rebalance"]["pools"][0]["triggered"] is True


def test_rebalance_buy_only_needs_new_cash(client):
    p = client.post("/api/workbook/rebalance", json={}).json()["pools"][0]
    assert p["mode"] == "buy_only" and p["orders"] == [] and p["total_value"] == 3_000_000
    p2 = client.post("/api/workbook/rebalance", json={"cash": {"me_wb_pool": {"amount": 1_000_000}}}).json()["pools"][0]
    assert p2["new_cash"] == 1_000_000 and p2["orders"] and all(o["side"] == "매수" for o in p2["orders"])
    assert sum(o["amount"] for o in p2["orders"]) <= 1_000_000
    assert client.post("/api/workbook/rebalance", json={"cash": {"me_wb_pool": {"amount": -1}}}).status_code == 422


def test_report_xlsx_starts_with_one_line_summary(client):
    from openpyxl import load_workbook
    r = client.get("/api/workbook/report.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["요약", "보유종목", "자산", "리밸런싱"]
    assert wb["요약"]["A1"].value.startswith(date.today().strftime("%Y.%m.%d") + " 기준 순자산")
    assert wb["보유종목"]["B2"].value == "KODEX 200"
