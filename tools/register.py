"""register — 처음 등록: 캡처 이미지 → 파싱 → 종목코드 채움 → CSV + 엑셀 저장.

예) uv run python -m tools.register --account me_mirae_pension_new --images a.png b.png c.png --out data/samples/x/holdings.csv
수량이 바뀔 때(거래·납입)만 다시 돌린다. 매번의 현재가는 tools.market.fetch_prices.
"""
from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

import pandas as pd

from tools.market import resolve_ticker
from tools.parse_balance import parse_capture

COLUMNS = ["account_id", "account_label", "ticker_name", "ticker", "qty", "buy_value", "price", "value"]


def register(account_id: str, images: list[Path], account_label: str | None = None) -> tuple[pd.DataFrame, dict]:
    df, meta = parse_capture(images, account_id)
    df["account_label"] = account_label or meta["account_label"]
    codes, pending = [], []
    for name in df["ticker_name"]:
        code, cands = resolve_ticker(name)
        codes.append(code)
        if code is None:
            pending.append((name, cands[:5]))
        time.sleep(0.2)
    df["ticker"] = codes
    meta["unresolved"] = pending
    return df[COLUMNS], meta


def save(df: pd.DataFrame, out_csv: Path) -> Path:
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_csv, index=False, encoding="utf-8-sig")
    out_xlsx = out_csv.with_suffix(".xlsx")
    df.to_excel(out_xlsx, index=False)
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(out_xlsx); ws = wb.active
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=5, max_col=6):
        for c in row:
            c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", [24, 18, 26, 10, 8, 12, 10, 12]):
        ws.column_dimensions[col].width = w
    wb.save(out_xlsx)
    return out_xlsx


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="캡처 → 보유 내역 등록")
    ap.add_argument("--account", required=True, action="append", help="계좌 id (여러 번 가능, --images와 짝)")
    ap.add_argument("--images", required=True, action="append", nargs="+", help="그 계좌의 캡처 파일들")
    ap.add_argument("--out", required=True, help="저장할 CSV 경로 (같은 이름의 .xlsx도 만든다)")
    args = ap.parse_args(argv)
    if len(args.account) != len(args.images):
        ap.error("--account 와 --images 개수가 달라요")
    frames = []
    for acct, imgs in zip(args.account, args.images):
        df, meta = register(acct, [Path(p) for p in imgs])
        ok = "일치" if meta["reconciled"] else f"불일치 {meta['diff']:+,}원 → 표를 확인·수정하세요"
        print(f"{acct}: {len(df)}종목, 화면 합계 {meta['screen_total_value']:,}원 대조 {ok}")
        for name, cands in meta["unresolved"]:
            print(f"  종목코드 미확정: {name} → 후보 {cands}")
        frames.append(df)
    out = Path(args.out)
    xlsx = save(pd.concat(frames, ignore_index=True), out)
    print(f"저장: {out}\n저장: {xlsx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
