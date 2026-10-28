"""workbook — 처음부터 엑셀처럼 보는 화면의 데이터 (Claude 담당, 2026-10-28 사용자 "처음부터 엑셀처럼 보여주면 안 됨?").

- /api/workbook/rebalance : 등록한 잔고(계좌별 최신)로 리밸런싱 묶음마다 기존 엔진(tools.rebalance.compute_rebalance)을 돌린다.
  이번 달 납입은 묶음별로 넣을 수 있다. 계산은 엔진, 문장은 tools.explain 템플릿 — LLM 없음.
- /api/workbook/summary   : 한 줄 요약과 요약 숫자 (대시보드 + 종목 잔고 + 리밸런싱 판정).
- /api/workbook/report.xlsx: 한 줄 요약이 맨 위에 있는 보고서 엑셀 (요약·보유종목·자산·리밸런싱 시트).
"""
from __future__ import annotations

import datetime as dt
import io
from dataclasses import asdict
from typing import Annotated, Literal

import pandas as pd
from fastapi import APIRouter
from fastapi.responses import Response
from pydantic import Field

from backend import holdings as hapi
from backend import registration, storage
from backend.storage import InputModel
from tools.assets import defaults_with_prices, load_accounts, load_asset_map, load_pools, tag_holdings
from tools.parse_balance import HOLDING_COLUMNS
from tools.strategy import CLASS_LABELS, load_strategy

router = APIRouter(prefix="/api/workbook", tags=["workbook"])
Owner = Literal["together", "me", "partner"]


class PoolCash(InputModel):
    amount: Annotated[int, Field(strict=True, ge=0, le=10**12)] = 0
    account_id: str | None = None


class RebalanceInput(InputModel):
    owner: Owner = "together"
    live: bool = False
    cash: dict[str, PoolCash] = {}


def _holdings_frame(owner: str, live: bool) -> tuple[pd.DataFrame, dict, list[str]]:
    accounts, rows = hapi.load_rows(owner)
    failed: list[str] = []
    prices = {}
    if live and rows:
        quotes, failed = hapi.quotes_for(sorted({h.ticker for h in rows if h.ticker}), hapi.settings()["quote_ttl_seconds"])
        prices = {k: v["price"] for k, v in quotes.items() if v.get("price")}
    df = pd.DataFrame([{"account_id": h.account_id, "ticker_name": h.ticker_name, "ticker": h.ticker, "qty": h.qty,
                        "price": prices.get(h.ticker or "", h.price)} for h in rows], columns=HOLDING_COLUMNS)
    as_of = {a.id: a.as_of.isoformat() if a.as_of else None for a in accounts}
    return df, as_of, failed


def rebalance_all(data: RebalanceInput) -> dict:
    from tools.explain import headline
    from tools.rebalance import NewCash, compute_rebalance
    accts = load_accounts(registration.accounts_path())
    pools = load_pools(registration.accounts_path())
    amap = load_asset_map(registration.asset_map_path())
    df, as_of, failed = _holdings_frame(data.owner, data.live)
    out = []
    for pid, pool in pools.items():
        ids = [a for a in pool.get("accounts", []) if a in accts and accts[a].get("active", True) is not False]
        part = df[df["account_id"].isin(ids)].reset_index(drop=True)
        if part.empty:
            continue
        label = pool.get("label") or pid
        base = {"pool_id": pid, "label": label, "mode": pool.get("mode", "rebalance"), "accounts": [{"id": a, "label": accts[a].get("label", a)} for a in ids],
                "as_of": min((as_of[a] for a in ids if as_of.get(a)), default=None)}
        try:
            strategy = load_strategy(pool.get("strategy") or "all_weather")
        except (KeyError, ValueError) as e:
            out.append({**base, "error": f"전략을 쓸 수 없어요: {e}"}); continue
        if not strategy.targets:
            out.append({**base, "error": "이 묶음의 전략 비율이 아직 정해지지 않았어요."}); continue
        tagged, unmapped = tag_holdings(part, amap)
        if unmapped:
            out.append({**base, "error": "자산군이 정해지지 않은 종목: " + ", ".join(unmapped)}); continue
        cash = data.cash.get(pid) or PoolCash()
        new_cash = NewCash(cash.amount, cash.account_id if cash.account_id in ids else ids[0]) if cash.amount else None
        try:
            defaults = defaults_with_prices(tagged, live=data.live, instruments=strategy.instruments, asset_map=amap)
            r = compute_rebalance(tagged, strategy, new_cash=new_cash, accounts=accts, mode=pool.get("mode", "rebalance"), defaults=defaults)
        except ValueError as e:
            out.append({**base, "error": str(e)}); continue
        res = asdict(r)
        out.append({**base, "strategy": r.strategy_name, "threshold_pp": round(r.threshold * 100, 2), "min_trade_krw": r.min_trade_krw,
                    "total_value": r.total_value, "new_cash": r.new_cash, "new_cash_account": new_cash.account_id if new_cash else ids[0],
                    "classes": res["classes"], "orders": res["orders"], "triggered": r.triggered, "blocked_reason": r.blocked_reason,
                    "constraints": res["constraints"], "notes": r.notes, "headline": headline(r)})
    return {"pools": out, "quote_failed": failed}


@router.post("/rebalance")
def rebalance(data: RebalanceInput):
    return rebalance_all(data)


def _won(n: int | None) -> str:
    return "—" if n is None else f"{n:,}원"


def _man(n: int) -> str:
    """1억 2,345만원 — 한 줄 요약용 (만원 아래 버림 없이 반올림)."""
    sign = "−" if n < 0 else ""
    n = abs(int(round(n / 10_000)))
    eok, man = divmod(n, 10_000)
    return sign + (f"{eok:,}억 " if eok else "") + (f"{man:,}만원" if man else ("원" if eok else "0원")).strip()


def summary(owner: str = "together", live: bool = False) -> dict:
    dash = storage.dashboard(owner)   # Codex 장부 집계(순자산·주머니)를 그대로 쓴다
    h = hapi.holdings(owner, live)
    rb = rebalance_all(RebalanceInput(owner=owner, live=live))
    t = h["total"]
    parts = [f"{dt.date.today():%Y.%m.%d} 기준 순자산 {_man(dash['net'])}"]
    if t["count"]:
        line = f"종목 잔고 {_man(t['value'])}"
        if t["total_return"] is not None:
            line += f"(넣은 돈 {_man(t['buy_value'])}, 불어난 돈 {_man(t['total_return'])} {t['total_rate']:+.2f}%)"
        parts.append(line)
    if live and t.get("today_return") is not None:
        parts.append(f"오늘 {_man(t['today_return'])}")
    need = []
    for p in rb["pools"]:
        if p.get("error") or not p.get("classes"):
            continue
        worst = max(p["classes"], key=lambda c: abs(c["deviation_pp"]))
        if p["triggered"]:
            need.append(f"{p['label'].split(' — ')[0]} {worst['label']} {worst['deviation_pp']:+.1f}%p")
    parts.append("리밸런싱 필요: " + ", ".join(need) if need else "리밸런싱: 모두 범위 안")
    stale = [a["label"] for a in h["accounts"] if a["stale"]]
    if stale:
        parts.append("다시 등록할 잔고: " + ", ".join(stale))
    return {"line": " · ".join(parts), "dashboard": dash, "holdings": h, "rebalance": rb}


@router.get("/summary")
def summary_api(owner: Owner = "together", live: bool = False):
    return summary(owner, live)


@router.get("/report.xlsx")
def report(owner: Owner = "together", live: bool = False):
    """한 줄 요약이 맨 위에 있는 보고서. 요약·보유종목·자산·리밸런싱 시트."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    s = summary(owner, live)
    d, h, rb = s["dashboard"], s["holdings"], s["rebalance"]
    wb = Workbook()
    head = PatternFill("solid", fgColor="F2F2F2")
    bold = Font(bold=True)

    def table(ws, header, rows, start_row):
        for j, v in enumerate(header, 1):
            c = ws.cell(start_row, j, v); c.font = bold; c.fill = head
        for i, row in enumerate(rows, start_row + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                if isinstance(v, int) and not isinstance(v, bool):
                    c.number_format = "#,##0"
                elif isinstance(v, float):
                    c.number_format = "0.00"
        return start_row + len(rows) + 2

    ws = wb.active; ws.title = "요약"
    ws["A1"] = s["line"]; ws["A1"].font = Font(bold=True, size=12); ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A1:H1"); ws.row_dimensions[1].height = 48
    r = table(ws, ["항목", "금액(원)"], [["순자산", d["net"]], ["총자산", d["total"]], ["대출", d["debt"]], ["노후 주머니", d["retirement"]],
                                        ["현금 주머니", d["cash"]], ["부동산", d["property"]]], 3)
    t = h["total"]
    table(ws, ["계좌", "기준일", "평가금액", "넣은 돈", "불어난 돈", "수익률(%)"],
          [[a["label"], a["as_of"], a["value"], a["buy_value"], a["total_return"], a["total_rate"]] for a in h["accounts"]]
          + [["합계", None, t["value"], t["buy_value"], t["total_return"], t["total_rate"]]], r)
    for col, w in zip("ABCDEFGH", [26, 14, 16, 16, 16, 12, 12, 12]):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("보유종목")
    labels = {a["id"]: a["label"] for a in h["accounts"]}
    table(ws2, ["계좌", "종목명", "종목코드", "보유수량", "현재가", "평가금액", "매입금액", "손익", "수익률(%)", "자산군", "비중(%)"],
          [[labels.get(x["account_id"], x["account_id"]), x["ticker_name"], x["ticker"], x["qty"], x["price"], x["value"], x["buy_value"],
            x["total_return"], x["total_rate"], CLASS_LABELS.get(x["asset_class"] or "", "미정"), x["weight"]] for x in h["holdings"]], 1)
    for col, w in zip("ABCDEFGHIJK", [20, 26, 10, 10, 10, 14, 14, 14, 10, 10, 8]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("자산")
    owner_l = {"me": "나", "partner": "배우자", "joint": "공동"}
    pocket_l = {"retirement": "노후", "cash": "현금", "property": "부동산", "debt": "대출"}
    table(ws3, ["이름", "계좌·보관처", "소유자", "종류", "자산군", "금액", "기준일"],
          [[a["name"], a["account"], owner_l.get(a["owner"], a["owner"]), pocket_l.get(a["pocket"], a["pocket"]), CLASS_LABELS.get(a["assetClass"], a["assetClass"]), a["amount"], a["as_of"]]
           for a in d["assets"] + d["debts"]], 1)
    for col, w in zip("ABCDEFG", [28, 24, 8, 8, 10, 16, 12]):
        ws3.column_dimensions[col].width = w

    ws4 = wb.create_sheet("리밸런싱")
    row = 1
    for p in rb["pools"]:
        ws4.cell(row, 1, f"{p['label']}" + (f" · {p['strategy']}" if p.get("strategy") else "")).font = bold
        row += 1
        if p.get("error"):
            ws4.cell(row, 1, p["error"]); row += 2; continue
        ws4.cell(row, 1, p["headline"]); row += 1
        row = table(ws4, ["자산군", "목표(%)", "현재(%)", "편차(%p)", "평가금액"],
                    [[c["label"], round(c["target"] * 100, 2), round(c["current"] * 100, 2), c["deviation_pp"], c["value"]] for c in p["classes"]], row)
        if p["orders"]:
            row = table(ws4, ["구분", "자산군", "종목", "수량", "금액", "비고"],
                        [[o["side"], CLASS_LABELS.get(o["asset_class"], o["asset_class"]), o["ticker_name"], o["qty"], o["amount"], o["note"]] for o in p["orders"]], row)
    for col, w in zip("ABCDEF", [30, 12, 26, 10, 16, 40]):
        ws4.column_dimensions[col].width = w

    buf = io.BytesIO(); wb.save(buf)
    fname = f"retire_report_{dt.date.today():%Y%m%d}.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"; filename*=UTF-8\'\'%EB%91%98%EC%9D%B4%EC%84%9C%EC%9D%80%ED%87%B4_%EB%B3%B4%EA%B3%A0%EC%84%9C_{dt.date.today():%Y%m%d}.xlsx'})
